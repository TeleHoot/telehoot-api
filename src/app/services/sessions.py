import enum
import io
import secrets
from typing import Any

import pandas as pd
from openpyxl.cell import MergedCell
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from src import core
from src.app import models, repositories, schemas, services


class Sessions(
    core.services.BaseCRUD[
        schemas.sessions.Create,
        schemas.sessions.Read,
        schemas.sessions.Update,
        schemas.sessions.Filters,
        schemas.sessions.SortParams,
        models.Session,
    ]
):
    def __init__(self):
        self.repo = repositories.Sessions()
        super().__init__(
            repo=self.repo,
            create_schema=schemas.sessions.Create,
            read_schema=schemas.sessions.Read,
            update_schema=schemas.sessions.Update,
            filters_schema=schemas.sessions.Filters,
        )

    async def create(
        self,
        uow: core.UnitOfWork,
        create_schema: schemas.sessions.Create,
        *,
        additional_data: dict[str, Any] | None = None,
    ) -> schemas.sessions.Read:
        max_retries = 10
        additional_data = additional_data or {}

        for _ in range(max_retries):
            code = str(secrets.SystemRandom().randrange(0, 10000)).zfill(4)

            exists = await uow.postgres_session.scalar(
                select(models.Session).filter_by(
                    join_code=code, status=models.SessionStatus.WAITING
                )
            )

            if exists:
                continue

            return await super().create(
                uow,
                create_schema,
                additional_data={
                    **additional_data,
                    "join_code": code,
                },
            )

        raise RuntimeError("Failed to generate unique join code after multiple attempts")

    async def get_session_results(
        self,
        uow: core.UnitOfWork,
        session: schemas.sessions.Read,
        *,
        include_answers: bool = False,
    ) -> list[schemas.sessions.SessionResult]:
        results = await self.repo.get_session_results(
            uow, session.id, include_answers=include_answers
        )

        if include_answers:
            participants_data = {}

            for row in results:
                participant: models.Participant = row.Participant
                total_points: int = row.total_points
                answer: models.ParticipantAnswer | None = row.ParticipantAnswer

                if (participant_id := str(participant.id)) not in participants_data:
                    participants_data[participant_id] = {
                        "participant": participant,
                        "total_points": total_points,
                        "answers": [],
                    }

                if answer:
                    participants_data[participant_id]["answers"].append(answer)

            return [
                schemas.sessions.SessionResult(
                    participant=data["participant"],
                    total_points=data["total_points"],
                    answers=data["answers"],
                )
                for data in participants_data.values()
            ]

        return [
            schemas.sessions.SessionResult(
                participant=row.Participant, total_points=row.total_points
            )
            for row in results
        ]

    async def export_session_results(
        self,
        uow: core.UnitOfWork,
        session: schemas.sessions.Read,
        questions_service: services.Questions,
    ) -> io.BytesIO:
        results: list[schemas.sessions.SessionResult] = await self.get_session_results(
            uow, session, include_answers=True
        )

        questions = await questions_service.read_many(
            uow,
            filters=schemas.questions.Filters(quiz_id=session.quiz.id),
            sorting=schemas.questions.SortParams(sort_by=schemas.questions.SortFields.CREATED_AT),
            pagination=core.schemas.PaginationParams(limit=100),
        )
        subheaders, group_headers = self._build_headers(questions)

        data_rows = self._build_data_rows(results, questions)

        df = pd.DataFrame(data_rows, columns=subheaders)  # type: ignore[reportGeneralTypeIssues]

        excel_file = io.BytesIO()
        with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:  # type: ignore[reportGeneralTypeIssues]
            self._write_dataframe_to_excel(writer, df, subheaders, group_headers)

            self._apply_styles_and_alignment(writer, df, group_headers)

        excel_file.seek(0)
        return excel_file

    @staticmethod
    def _build_headers(
        questions: list[schemas.questions.Read],
    ) -> tuple[list[str], list[tuple[str, int, int]]]:
        subheaders: list[str] = []

        user_info_columns = [
            Labels.TELEGRAM_LOGIN,
            Labels.SESSION_USERNAME,
            Labels.POINTS,
            Labels.CORRECT_ANSWERS,
            Labels.INCORRECT_ANSWERS,
        ]
        subheaders.extend(user_info_columns)

        group_headers: list[tuple[str, int, int]] = [
            (Labels.USER_INFO, 0, len(user_info_columns) - 1)
        ]

        for i, _ in enumerate(questions):
            group_start = len(subheaders)
            subheaders.extend([
                f"Вопрос {i + 1} - {Labels.QUESTION}",
                f"Вопрос {i + 1} - {Labels.ANSWERS}",
                f"Вопрос {i + 1} - {Labels.IS_CORRECT}",
            ])
            group_headers.append((f"Вопрос {i + 1}", group_start, group_start + 2))

        subheaders.insert(0, "")
        group_headers = [(name, start + 1, end + 1) for (name, start, end) in group_headers]

        return subheaders, group_headers

    @staticmethod
    def _build_data_rows(
        results: list[schemas.sessions.SessionResult], questions: list[schemas.questions.Read]
    ) -> list[dict[str, Any]]:
        data_rows = []
        for result in results:
            participant = result.participant
            answers = {a.question_id: a for a in result.answers}

            correct_answers = sum(1 for a in result.answers if a.is_correct)
            incorrect_answers = len(result.answers) - correct_answers

            row: dict[str, Any] = {
                Labels.TELEGRAM_LOGIN: participant.user.telegram_username,
                Labels.SESSION_USERNAME: participant.session_nickname,
                Labels.POINTS: result.total_points,
                Labels.CORRECT_ANSWERS: correct_answers,
                Labels.INCORRECT_ANSWERS: incorrect_answers,
            }

            for i, question in enumerate(questions, 1):
                answer = answers.get(question.id)
                row.update({
                    f"Вопрос {i} - {Labels.QUESTION}": question.title,
                    f"Вопрос {i} - {Labels.ANSWERS}": answer.text if answer else "-",
                    f"Вопрос {i} - {Labels.IS_CORRECT}": answer.is_correct if answer else "-",
                })

            data_rows.append(row)

        return data_rows

    @staticmethod
    def _write_dataframe_to_excel(
        writer: pd.ExcelWriter,
        df: pd.DataFrame,
        subheaders: list[str],
        group_headers: list[tuple[str, int, int]],
    ) -> None:
        df.to_excel(
            writer,
            index=False,
            startrow=ExcelPositions.SUBHEADER_ROW,
            header=False,
            sheet_name=Labels.SHEET_NAME,
        )
        worksheet = writer.sheets[Labels.SHEET_NAME]

        for group_name, start_col, end_col in group_headers:
            worksheet.merge_cells(
                start_row=ExcelPositions.GROUP_HEADER_ROW,
                end_row=ExcelPositions.GROUP_HEADER_ROW,
                start_column=start_col + 1,
                end_column=end_col + 1,
            )
            worksheet.cell(ExcelPositions.GROUP_HEADER_ROW, start_col + 1, group_name)

        for col_num, value in enumerate(subheaders, 1):
            worksheet.cell(ExcelPositions.SUBHEADER_ROW, col_num, value.split(" - ")[-1])

        for idx in range(1, len(df.columns) + 1):
            column_letter = get_column_letter(idx)
            cells = [cell for cell in worksheet[column_letter] if not isinstance(cell, MergedCell)]
            if cells:
                max_length = max(len(str(cell.value)) for cell in cells)
                worksheet.column_dimensions[column_letter].width = max_length + 2

    @staticmethod
    def _apply_styles_and_alignment(
        writer: pd.ExcelWriter,
        df: pd.DataFrame,
        group_headers: list[tuple[str, int, int]],
    ) -> None:
        worksheet = writer.sheets[Labels.SHEET_NAME]

        medium_side = Side(border_style="medium", color="000000")
        thin_side = Side(border_style="thin", color="000000")

        first_row = ExcelPositions.GROUP_HEADER_ROW
        last_row = ExcelPositions.SUBHEADER_ROW + len(df)
        first_col = ExcelPositions.EMPTY_COLUMN
        last_col = len(df.columns)

        section_border_columns = set()
        for _, _start_col, end_col in group_headers[:-1]:
            section_border_columns.add(end_col + 1)

        for xl_row in range(first_row, last_row + 1):
            for col in range(first_col, last_col + 1):
                cell = worksheet.cell(xl_row, col)

                if col == ExcelPositions.EMPTY_COLUMN:
                    left = None
                    top = None
                    bottom = None
                    right = thin_side
                elif col == ExcelPositions.FIRST_DATA_COLUMN:
                    left = medium_side
                    top = medium_side if xl_row == first_row else thin_side
                    bottom = medium_side if xl_row == last_row else thin_side
                    right = thin_side
                else:
                    left = thin_side
                    top = medium_side if xl_row == first_row else thin_side
                    bottom = medium_side if xl_row == last_row else thin_side

                    if col == last_col or col in section_border_columns:
                        right = medium_side
                    else:
                        right = thin_side

                cell.border = Border(
                    left=left if left else Side(border_style=None),
                    right=right if right else Side(border_style=None),
                    top=top if top else Side(border_style=None),
                    bottom=bottom if bottom else Side(border_style=None),
                )

        center_alignment = Alignment(horizontal="center", vertical="center")

        for xl_row in range(first_row, last_row + 1):
            for col in range(first_col, last_col + 1):
                cell = worksheet.cell(xl_row, col)
                cell.alignment = center_alignment


class Labels(enum.StrEnum):
    TELEGRAM_LOGIN = "Телеграм Логин"
    SESSION_USERNAME = "Имя в сессии"
    POINTS = "Очки"
    CORRECT_ANSWERS = "Правильные ответы"
    INCORRECT_ANSWERS = "Неправильные ответы"
    QUESTION = "Вопрос"
    ANSWERS = "Ответ"
    IS_CORRECT = "Правильный"
    USER_INFO = "Информация о пользователе"
    SHEET_NAME = "Лист1"


class ExcelPositions(enum.IntEnum):
    EMPTY_COLUMN = 1
    FIRST_DATA_COLUMN = 2
    GROUP_HEADER_ROW = 2
    SUBHEADER_ROW = 3
